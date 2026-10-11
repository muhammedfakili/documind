"""Format-specific loaders: file on disk -> list of Documents.

Every loader writes its text in the same light Markdown dialect -- "## " for a
heading, one line per table row with the column names kept -- so the chunker in
A3 can split all formats the same way.

Granularity differs on purpose:
  Markdown  one Document per file   (chunking splits it later, by heading)
  PDF       one Document per page   (so a citation can say "page 3")
  CSV/Excel one Document per row    (a row is already a self-contained fact)
"""

from __future__ import annotations

import csv
from collections.abc import Callable
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook

from app.ingest.documents import Document, classify, extract_effective_date, parse_iso_date

# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------


def load_markdown(path: Path, rel: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8-sig").strip()
    title = next((ln[2:].strip() for ln in text.splitlines() if ln.startswith("# ")), path.stem)
    doc_type, audience = classify(rel, text)
    return [
        Document(
            text=text,
            source=rel.as_posix(),
            title=title,
            doc_type=doc_type,
            audience=audience,
            effective_date=extract_effective_date(text),
        )
    ]


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def table_to_lines(rows: list[list[str | None]]) -> list[str]:
    """Render a table one row per line, repeating the column names.

    Plain text extraction reads a table left to right, line by line, so a
    wrapped cell ends up split across lines and "E-214" lands far from its
    meaning. Keeping "Column: value" on every row means each line still makes
    sense on its own -- even after the chunker cuts the table in half.
    """
    cells = [[" ".join((c or "").split()) for c in row] for row in rows if any(row)]
    if len(cells) < 2:
        return [" | ".join(r) for r in cells]
    header, body = cells[0], cells[1:]
    if len(header) == 2:
        # Two-column tables are key/value lists ("Driver | 10 mm dynamic").
        return [f"{key}: {value}" for key, value in body]
    return [" | ".join(f"{h}: {v}" for h, v in zip(header, row) if v) for row in body]


# Measured on our manuals: wrapped lines of one paragraph are 4.2 pt apart,
# separate paragraphs 7 pt or more. Other PDF generators need re-measuring.
_PARAGRAPH_GAP = 5.0


def _is_heading(line: dict) -> bool:
    chars = [c for c in line.get("chars", []) if c["text"].strip()]
    return bool(chars) and all("Bold" in c["fontname"] for c in chars)


def pdf_page_text(page: pdfplumber.page.Page) -> str:
    """Text of one page with tables rendered as rows and headings marked."""
    tables = page.find_tables()
    rest = page
    for table in tables:
        rest = rest.outside_bbox(table.bbox)

    # A PDF has no paragraphs, only positioned lines. Lines closer together than
    # _PARAGRAPH_GAP belong to the same paragraph and are joined back into one
    # line; otherwise "blinks amber 3" and "times." would be separate lines.
    items: list[tuple[float, str]] = []
    prev_bottom: float | None = None
    for line in rest.extract_text_lines():
        text = line["text"].strip()
        if not text:
            continue
        if _is_heading(line):
            items.append((line["top"], f"## {text}"))
            prev_bottom = None
            continue
        continues = prev_bottom is not None and line["top"] - prev_bottom < _PARAGRAPH_GAP
        if continues and items and not items[-1][1].startswith("## "):
            top, previous = items[-1]
            items[-1] = (top, f"{previous} {text}")
        else:
            items.append((line["top"], text))
        prev_bottom = line["bottom"]

    # Interleave text and tables by vertical position, so each table stays
    # under the heading it belongs to.
    for table in tables:
        items.append((table.bbox[1], "\n".join(table_to_lines(table.extract()))))
    items.sort(key=lambda item: item[0])
    return "\n".join(text for _, text in items)


def load_pdf(path: Path, rel: Path) -> list[Document]:
    with pdfplumber.open(path) as pdf:
        pages = [pdf_page_text(p) for p in pdf.pages]
        title = (pdf.metadata or {}).get("Title") or path.stem
    full_text = "\n".join(pages)
    doc_type, audience = classify(rel, full_text)
    effective = extract_effective_date(full_text)  # stated once, applies to every page
    return [
        Document(
            text=text,
            source=rel.as_posix(),
            title=title,
            doc_type=doc_type,
            audience=audience,
            effective_date=effective,
            page=number,
        )
        for number, text in enumerate(pages, start=1)
        if text.strip()
    ]


# ---------------------------------------------------------------------------
# Row-based formats: CSV and Excel
# ---------------------------------------------------------------------------
# A row is turned into a sentence-like line that keeps every column name. An
# embedding of "12 | 24 | No" means nothing; "Warranty US region (months): 12"
# is something a question can match.

RowFormatter = Callable[[dict[str, str]], str]


def generic_row(row: dict[str, str]) -> str:
    return "; ".join(f"{key}: {value}" for key, value in row.items() if value not in (None, ""))


def ticket_row(row: dict[str, str]) -> str:
    return (
        f"Support ticket {row['ticket_id']} ({row['created_at']}, {row['channel']}, "
        f"country {row['country']}, status {row['status']}). "
        f"Product: {row['product']}. Category: {row['category']}.\n"
        f"Subject: {row['subject']}\n"
        f"Customer wrote: {row['customer_message']}\n"
        f"Resolution: {row['resolution']}"
    )

COUNTRIES = {"DE": "Germany", "UK": "United Kingdom", "US": "United States",
             "NO": "Norway", "FR": "France", "NL": "Netherlands", "AU": "Australia"}


def partner_row(row: dict[str, str]) -> str:
    kind = "physical store" if row["type"] == "store" else "online shop"
    country = COUNTRIES[row["country"]]
    products = row["carries"].replace("; ", ", ")
    return (
        f"Retail partner {row['partner']}: {kind} in {row['city']}, {country}. "
        f"Partner tier: {row['tier']}.\n"
        f"Carries: {products}. Demo units available in store: {row['demo_units']}."
    )


# Filename -> formatter. Anything not listed falls back to generic_row.
CSV_FORMATTERS: dict[str, RowFormatter] = {
    "tickets-2026-q3.csv": ticket_row,
    "retail-partners.csv": partner_row,
}


def load_csv(path: Path, rel: Path) -> list[Document]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    formatter = CSV_FORMATTERS.get(path.name, generic_row)
    doc_type, audience = classify(rel, "")
    title = path.stem.replace("-", " ").capitalize()
    return [
        Document(
            text=formatter(row),
            source=rel.as_posix(),
            title=title,
            doc_type=doc_type,
            audience=audience,
            # A ticket's own date is when it was created.
            effective_date=parse_iso_date(row.get("created_at", "")),
            row=number,
        )
        for number, row in enumerate(rows, start=1)
    ]


def load_xlsx(path: Path, rel: Path) -> list[Document]:
    wb = load_workbook(path, read_only=True, data_only=True)
    doc_type, audience = classify(rel, "")
    docs: list[Document] = []
    for ws in wb.worksheets:
        rows = [r for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
        if len(rows) < 2:
            continue
        header = [str(h) for h in rows[0]]
        title = f"{path.stem.replace('-', ' ').capitalize()} / {ws.title}"
        for number, values in enumerate(rows[1:], start=1):
            row = {h: ("" if v is None else str(v)) for h, v in zip(header, values)}
            docs.append(
                Document(
                    # The sheet name is context the row itself doesn't carry.
                    text=f"{title}. {generic_row(row)}",
                    source=rel.as_posix(),
                    title=title,
                    doc_type=doc_type,
                    audience=audience,
                    row=number,
                )
            )
    wb.close()
    return docs
